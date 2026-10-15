from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from apps.auditoria.models import Hallazgo


@login_required
def exportar_excel(request):
    """CU5 (cierre): exporta el listado de hallazgos con su evidencia asociada."""
    qs = Hallazgo.objects.select_related(
        "indicador_principal", "indicador_relacionado", "regla_aplicada"
    ).order_by("-creado_en")

    tipo = request.GET.get("tipo")
    estado = request.GET.get("estado")
    dependencia = request.GET.get("dependencia")
    if tipo:
        qs = qs.filter(tipo=tipo)
    if estado:
        qs = qs.filter(estado=estado)
    if dependencia:
        qs = qs.filter(indicador_principal__dependencia_id=dependencia)
    ejecucion = request.GET.get("ejecucion")
    if ejecucion and ejecucion.isdigit():
        qs = qs.filter(ejecucion_id=ejecucion)

    wb = Workbook()
    ws = wb.active
    ws.title = "Hallazgos"

    encabezados = [
        "ID", "Tipo", "Estado", "Indicador principal", "Dependencia principal",
        "Indicador relacionado", "Score similitud", "Regla aplicada",
        "Parámetro de gobernanza", "Evidencia", "Revisado por", "Observación de revisión", "Creado en",
    ]
    ws.append(encabezados)

    for h in qs:
        ws.append([
            h.pk,
            h.get_tipo_display(),
            h.get_estado_display(),
            h.indicador_principal.nombre,
            str(h.indicador_principal.dependencia),
            h.indicador_relacionado.nombre if h.indicador_relacionado else "",
            h.score_similitud,
            h.regla_aplicada.nombre,
            h.regla_aplicada.parametro_gobernanza.nombre,
            h.evidencia,
            str(h.revisado_por) if h.revisado_por else "",
            h.observacion_revision,
            h.creado_en.strftime("%Y-%m-%d %H:%M"),
        ])

    for i, _ in enumerate(encabezados, start=1):
        ws.column_dimensions[get_column_letter(i)].width = 28

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = 'attachment; filename="hallazgos_auditoria.xlsx"'
    wb.save(response)
    return response
