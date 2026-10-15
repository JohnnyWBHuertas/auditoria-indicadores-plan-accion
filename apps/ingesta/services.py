"""Servicio de ingesta (CU1): carga un archivo y normaliza la hoja de vida.

Decision de arquitectura: el cargue se resuelve por PlantillaMapeoColumnas
(columna_origen -> campo_hoja_de_vida), no por posicion fija de columna,
porque la estructura del archivo exportado puede cambiar entre vigencias
y el prototipo se disena desacoplado de un dataset de referencia fijo.
"""

from __future__ import annotations

import csv
import io
import os
import re

from django.conf import settings
from django.db import transaction
from openpyxl import load_workbook

from apps.gobernanza.models import Dependencia

from .models import CargaArchivo, HistorialCambioIndicador, Indicador, LineaPAA

CAMPOS_HOJA_DE_VIDA = [
    "objetivo", "formula", "unidad_medida", "linea_base",
    "meta", "periodicidad", "fuente", "responsable",
]


def _leer_tabla(archivo_field) -> tuple[list[str], list[dict]]:
    """Lee un CSV o un Excel (.xlsx) y devuelve (columnas, filas). Todos los
    valores se tratan como texto: la normalizacion posterior decide el tipo."""
    nombre = archivo_field.name.lower()
    archivo_field.open("rb")
    try:
        contenido = archivo_field.read()
    finally:
        archivo_field.close()

    if nombre.endswith(".csv"):
        texto = contenido.decode("utf-8-sig")
        try:
            dialecto = csv.Sniffer().sniff(texto[:4096], delimiters=",;\t")
        except csv.Error:
            dialecto = csv.excel
        lector = csv.DictReader(io.StringIO(texto), dialect=dialecto)
        columnas = [c.strip() for c in (lector.fieldnames or [])]
        filas = [{(k or "").strip(): (v or "") for k, v in fila.items()} for fila in lector]
        return columnas, filas

    libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    hoja = libro.worksheets[0]
    registros = list(hoja.iter_rows(values_only=True))
    if not registros:
        return [], []
    columnas = [str(c).strip() if c is not None else "" for c in registros[0]]
    filas = [
        {columnas[i]: ("" if valor is None else str(valor)) for i, valor in enumerate(fila) if i < len(columnas)}
        for fila in registros[1:]
        if any(valor not in (None, "") for valor in fila)
    ]
    return columnas, filas


def procesar_archivo(carga: CargaArchivo) -> None:
    """Procesa la carga y la deja siempre en un estado final: si algo falla de
    forma inesperada queda 'fallido' con el detalle, nunca 'procesando'."""
    if not carga.nombre_archivo and carga.archivo:
        carga.nombre_archivo = os.path.basename(carga.archivo.name)
        carga.save(update_fields=["nombre_archivo"])
    try:
        _procesar(carga)
    except Exception as exc:
        carga.estado = "fallido"
        carga.detalle_errores = [f"Error inesperado al procesar el archivo: {exc}"]
        carga.save(update_fields=["estado", "detalle_errores"])
    finally:
        if not settings.CONSERVAR_ARCHIVO_ORIGINAL:
            descartar_original(carga)


def descartar_original(carga: CargaArchivo) -> None:
    """Elimina del disco el archivo cargado; el registro de la carga se conserva."""
    if carga.archivo:
        carga.archivo.delete(save=False)
    CargaArchivo.objects.filter(pk=carga.pk).update(archivo="")
    carga.archivo = ""


def _procesar(carga: CargaArchivo) -> None:
    mapeo = carga.plantilla.mapeo_columnas  # {"columna_origen": "campo_hoja_de_vida"}

    try:
        columnas, filas = _leer_tabla(carga.archivo)
    except Exception as exc:
        carga.estado = "fallido"
        carga.detalle_errores = [f"No fue posible leer el archivo: {exc}"]
        carga.save(update_fields=["estado", "detalle_errores"])
        return

    columnas_mapeadas = {origen: campo for origen, campo in mapeo.items() if origen in columnas}
    errores = []
    faltantes = [c for c in mapeo if c not in columnas]
    if faltantes:
        errores.append(f"Columnas de la plantilla no encontradas en el archivo: {faltantes}")

    carga.filas_totales = len(filas)
    filas_cargadas = 0
    filas_con_error = 0
    filas_sin_indicador = 0

    # Una sola transaccion para todo el archivo (mucho mas rapido que confirmar
    # cada fila) y un punto de guardado por fila: un error en una fila no
    # descarta las demas.
    with transaction.atomic():
        for indice_fila, fila in enumerate(filas):
            try:
                datos = {campo: str(fila.get(origen, "")).strip() for origen, campo in columnas_mapeadas.items()}
                if not datos.get("nombre"):
                    # Linea del PAA sin indicador asociado: no es un error de carga.
                    filas_sin_indicador += 1
                    continue
                with transaction.atomic():
                    _normalizar_fila(carga, datos)
                filas_cargadas += 1
            except Exception as exc:
                filas_con_error += 1
                errores.append(f"Fila {indice_fila + 2}: {exc}")

    carga.filas_cargadas = filas_cargadas
    carga.filas_con_error = filas_con_error
    carga.filas_sin_indicador = filas_sin_indicador
    carga.detalle_errores = errores[:50]
    carga.estado = "completado" if not errores else "completado_con_errores"
    carga.save(
        update_fields=[
            "filas_totales", "filas_cargadas", "filas_con_error", "filas_sin_indicador", "detalle_errores", "estado",
        ]
    )


def _normalizar_periodicidad(valor: str) -> str:
    """Lleva la periodicidad del archivo a las opciones de la hoja de vida
    para que 'Anual' y 'anual' no se interpreten como una divergencia."""
    opciones = {clave for clave, _ in Indicador.PERIODICIDAD_CHOICES}
    limpio = valor.strip().lower().replace(" ", "_")
    return limpio if limpio in opciones else "no_especificada"


def _normalizar_vigencia(valor: str, por_defecto: str) -> str:
    """Extrae el anio de valores como 'Año 2025'; si no hay anio, usa el de la carga."""
    coincidencia = re.search(r"\b(20\d{2})\b", valor or "")
    return coincidencia.group(1) if coincidencia else por_defecto


def _normalizar_fila(carga: CargaArchivo, datos: dict) -> None:
    nombre_dependencia = datos.get("dependencia", "").strip() or "Sin dependencia"
    dependencia, _ = Dependencia.objects.get_or_create(nombre=nombre_dependencia)

    vigencia = _normalizar_vigencia(datos.get("vigencia", ""), carga.vigencia)

    clave = dict(
        dependencia=dependencia,
        programa=datos.get("programa", "").strip(),
        bien_servicio=datos.get("bien_servicio", "").strip(),
        nombre=datos.get("nombre", "").strip(),
        producto_actividad_proyecto=datos.get("producto_actividad_proyecto", "").strip(),
        vigencia=vigencia,
    )
    if not clave["nombre"]:
        raise ValueError("La fila no tiene nombre de indicador (columna mapeada a 'nombre').")

    campos_hoja_de_vida = {campo: datos.get(campo, "").strip() for campo in CAMPOS_HOJA_DE_VIDA}
    if campos_hoja_de_vida["periodicidad"]:
        campos_hoja_de_vida["periodicidad"] = _normalizar_periodicidad(campos_hoja_de_vida["periodicidad"])

    indicador, creado = Indicador.objects.get_or_create(
        **clave, defaults={**campos_hoja_de_vida, "carga_origen": carga}
    )

    if not creado:
        _registrar_cambios(indicador, campos_hoja_de_vida, carga)

    codigo_paa = datos.get("codigo_paa", "").strip()
    if codigo_paa:
        LineaPAA.objects.get_or_create(
            indicador=indicador, codigo_paa=codigo_paa, defaults={"vigencia": vigencia}
        )


def _registrar_cambios(indicador: Indicador, campos_nuevos: dict, carga: CargaArchivo) -> None:
    """Historial append-only: solo se registra cuando el valor nuevo difiere
    y no está vacío (una carga posterior no debe borrar un dato ya bueno)."""
    cambios = []
    for campo, valor_nuevo in campos_nuevos.items():
        valor_anterior = getattr(indicador, campo)
        if valor_nuevo and valor_nuevo != valor_anterior:
            cambios.append(
                HistorialCambioIndicador(
                    indicador=indicador,
                    campo=campo,
                    valor_anterior=valor_anterior,
                    valor_nuevo=valor_nuevo,
                    origen="carga_archivo",
                )
            )
            setattr(indicador, campo, valor_nuevo)

    if cambios:
        HistorialCambioIndicador.objects.bulk_create(cambios)
        indicador.carga_origen = carga
        indicador.save()
