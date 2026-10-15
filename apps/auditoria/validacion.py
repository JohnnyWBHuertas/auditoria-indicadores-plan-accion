"""Validacion funcional del prototipo sobre el conjunto de control.

Ejecuta los casos de prueba de la matriz de validacion funcional (Anexo D):
carga el conjunto de control de nueve indicadores, corre la auditoria y
contrasta, para cada caso, el resultado esperado con el obtenido. Todo
ocurre dentro de una transaccion que se revierte al terminar, de modo que no
deja datos en la base ni archivos cargados.

Por defecto usa el modelo de embeddings real. Con usar_modelo=False los pares
candidatos se simulan (los tres pares que el diseno del conjunto de control
debe producir), lo que permite ejecutar la matriz sin cargar el modelo.
"""

from __future__ import annotations

import csv
import io
import socket
import tempfile
from contextlib import contextmanager
from pathlib import Path
from unittest import mock

from django.conf import settings
from django.contrib.auth import get_user_model
from django.contrib.auth.models import AnonymousUser
from django.core.files.base import ContentFile
from django.core.management import call_command
from django.db import transaction
from django.test import RequestFactory, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.ingesta.models import CargaArchivo, HistorialCambioIndicador, Indicador, LineaPAA, PlantillaMapeoColumnas
from apps.ingesta.services import procesar_archivo
from apps.nlp_engine import embeddings as nlp
from apps.nlp_engine.models import ParCandidato

from .models import Hallazgo
from .reglas import ejecutar_auditoria, es_trazable

RUTA_CSV = Path(settings.BASE_DIR) / "datos_ejemplo" / "conjunto_control_2026.csv"
VIGENCIA = "2026"

# Pares que el diseno del conjunto de control debe producir: (nombre, producto) de cada indicador.
PARES_ESPERADOS = [
    (("Proyectos culturales apoyados mediante el Programa Nacional de Concertación Cultural", "Modalidad Artes Visuales"),
     ("Proyectos culturales apoyados mediante el Programa Nacional de Concertación Cultural", "Modalidad Música")),
    (("Número de acciones de circulación artística y cultural realizadas", "Agenda expositiva"),
     ("Actividades de divulgación artística y cultural desarrolladas", "Agenda pedagógica")),
    (("Bibliotecas públicas asistidas técnicamente", "Visitas territoriales"),
     ("Bibliotecas públicas asistidas técnicamente", "Convocatoria anual")),
]
INDICADOR_LIMPIO = "Porcentaje de trámites de expedición de CDP tramitados dentro del término"


class _Revertir(Exception):
    pass


def _indicador(nombre, producto):
    return Indicador.objects.get(nombre=nombre, producto_actividad_proyecto=producto, vigencia=VIGENCIA)


def _pares_simulados(vigencia, umbral=None):
    for a, b in PARES_ESPERADOS:
        ia, ib = sorted([_indicador(*a), _indicador(*b)], key=lambda i: i.pk)
        ParCandidato.objects.get_or_create(
            indicador_a=ia, indicador_b=ib, modelo=settings.NLP_ENGINE["EMBEDDING_MODEL"],
            defaults={"similitud_coseno": 0.9},
        )
    return len(PARES_ESPERADOS)


@contextmanager
def _sin_red():
    """Falla ante cualquier intento de conexion saliente."""
    def bloqueado(*args, **kwargs):
        raise RuntimeError("Intento de conexión externa durante la auditoría")
    with mock.patch.object(socket.socket, "connect", bloqueado), mock.patch.object(socket, "create_connection", bloqueado):
        yield


def cargar_conjunto_control():
    """CU1: carga el CSV del conjunto de control con su plantilla de mapeo."""
    call_command("seed_plantilla_demo", stdout=io.StringIO())
    plantilla = PlantillaMapeoColumnas.objects.get(nombre="Conjunto de control 2026")
    carga = CargaArchivo(plantilla=plantilla, vigencia=VIGENCIA)
    carga.archivo.save("conjunto_control_2026.csv", ContentFile(RUTA_CSV.read_bytes()), save=False)
    carga.save()
    procesar_archivo(carga)
    carga.refresh_from_db()
    return carga


def _tipos(ejecucion):
    return sorted(ejecucion.hallazgos.values_list("tipo", flat=True))


def _fila(codigo, modulo, regla, entrada, esperado, obtenido, cumple):
    return {
        "caso": codigo, "modulo": modulo, "regla": regla, "entrada": entrada, "esperado": esperado,
        "obtenido": obtenido, "estado": "Cumple" if cumple else "No cumple",
    }


def ejecutar_matriz(usar_modelo: bool = True, aislada: bool = False) -> list[dict]:
    """Devuelve las filas de la matriz de validacion funcional. Con aislada=True
    corre sobre una base de datos temporal (como las pruebas automatizadas), de
    modo que los datos reales cargados en la base del proyecto no la afecten."""
    if aislada:
        from django.test.runner import DiscoverRunner

        runner = DiscoverRunner(verbosity=0, interactive=False)
        configuracion = runner.setup_databases()
        try:
            return ejecutar_matriz(usar_modelo=usar_modelo, aislada=False)
        finally:
            runner.teardown_databases(configuracion)

    filas: list[dict] = []
    with tempfile.TemporaryDirectory() as tmp, override_settings(MEDIA_ROOT=tmp):
        try:
            with transaction.atomic():
                filas.extend(_casos(usar_modelo))
                raise _Revertir
        except _Revertir:
            pass
    return filas


def _casos(usar_modelo: bool) -> list[dict]:
    filas = []
    call_command("seed_catalogo_gobernanza", stdout=io.StringIO())
    call_command("seed_roles", stdout=io.StringIO())

    # --- CU1: ingesta y normalizacion ---------------------------------------
    carga = cargar_conjunto_control()
    filas.append(_fila(
        "CP-01", "CU1 Ingesta", "Plantilla de mapeo de columnas",
        "CSV de 9 indicadores y plantilla 'Conjunto de control 2026'",
        "Estado completado, 9 filas cargadas, 0 con error",
        f"Estado {carga.estado}, {carga.filas_cargadas} cargadas, {carga.filas_con_error} con error",
        carga.estado == "completado" and carga.filas_cargadas == 9 and carga.filas_con_error == 0,
    ))

    with open(RUTA_CSV, encoding="utf-8-sig", newline="") as f:
        filas_csv = list(csv.DictReader(f))
    primera = filas_csv[0]
    ind = _indicador(primera["Indicador"], primera["Producto, actividad o proyecto"])
    preservados = ind.formula == primera["Fórmula"] and ind.meta == primera["Meta"] and ind.unidad_medida == primera["Unidad de medida"]
    filas.append(_fila(
        "CP-02", "CU1 Ingesta", "Preservación de la estructura de la hoja de vida",
        "Primera fila del CSV", "Fórmula, meta y unidad de medida idénticas a las del archivo; una línea PAA por indicador",
        f"Atributos preservados: {preservados}; indicadores {Indicador.objects.filter(vigencia=VIGENCIA).count()}, "
        f"líneas PAA {LineaPAA.objects.filter(vigencia=VIGENCIA).count()}",
        preservados and Indicador.objects.filter(vigencia=VIGENCIA).count() == 9
        and LineaPAA.objects.filter(vigencia=VIGENCIA).count() == 9,
    ))

    antes_ind, antes_hist = Indicador.objects.filter(vigencia=VIGENCIA).count(), HistorialCambioIndicador.objects.count()
    segunda = cargar_conjunto_control()
    despues_ind = Indicador.objects.filter(vigencia=VIGENCIA).count()
    filas.append(_fila(
        "CP-03", "CU1 Ingesta", "Historial de solo adición",
        "Recarga del mismo archivo", "Sin indicadores nuevos ni cambios en el historial",
        f"Indicadores {antes_ind}→{despues_ind}; historial {antes_hist}→{HistorialCambioIndicador.objects.count()}",
        despues_ind == antes_ind and HistorialCambioIndicador.objects.count() == antes_hist and segunda.estado == "completado",
    ))

    # --- CU2-CU4: auditoria -------------------------------------------------
    usuario = get_user_model().objects.create_user("validador", password=None)
    if usar_modelo:
        ejecucion = ejecutar_auditoria(VIGENCIA, usuario=usuario)
    else:
        with mock.patch.object(nlp, "generar_embeddings_pendientes", lambda: 0), \
                mock.patch.object(nlp, "generar_pares_candidatos", _pares_simulados), _sin_red():
            ejecucion = ejecutar_auditoria(VIGENCIA, usuario=usuario)

    hallazgos = list(ejecucion.hallazgos.select_related("indicador_principal", "indicador_relacionado"))
    por_tipo = {t: [h for h in hallazgos if h.tipo == t] for t in
                ("clonacion", "duplicidad", "divergencia_operativa", "incompletitud_atributos", "problema_metodologico")}

    pares = ParCandidato.objects.filter(indicador_a__vigencia=VIGENCIA)
    pares_ids = {frozenset((p.indicador_a_id, p.indicador_b_id)) for p in pares}
    esperados = {frozenset((_indicador(*a).pk, _indicador(*b).pk)) for a, b in PARES_ESPERADOS}
    limpio = Indicador.objects.get(nombre=INDICADOR_LIMPIO)
    filas.append(_fila(
        "CP-04", "CU2 Comparación semántica", "Pares candidatos (similitud ≥ umbral)",
        "9 indicadores; umbral de duplicidad", "Los 3 pares diseñados y ninguno con el indicador sin inconsistencia",
        f"{len(pares_ids)} pares; los esperados presentes: {esperados <= pares_ids}; limpio en algún par: {any(limpio.pk in p for p in pares_ids)}",
        esperados <= pares_ids and not any(limpio.pk in p for p in pares_ids),
    ))

    filas.append(_fila(
        "CP-05", "CU3 Clonación", "Coincidencia textual casi exacta con variación marginal",
        "Dos indicadores con igual nombre, fórmula y meta y distinto producto/actividad", "1 hallazgo de clonación",
        f"{len(por_tipo['clonacion'])} hallazgo(s)", len(por_tipo["clonacion"]) == 1,
    ))
    filas.append(_fila(
        "CP-06", "CU2 Duplicidad", "Similitud semántica sobre umbral",
        "Dos indicadores con propósito equivalente y redacción distinta", "1 hallazgo de duplicidad",
        f"{len(por_tipo['duplicidad'])} hallazgo(s)", len(por_tipo["duplicidad"]) == 1,
    ))
    div = por_tipo["divergencia_operativa"]
    campos_div = sorted(div[0].atributos_comparados["indicador_a"]) if div else []
    filas.append(_fila(
        "CP-07", "CU4 Divergencia operativa", "Equivalencia semántica con unidad o periodicidad divergentes",
        "Mismo nombre; unidad Número/Porcentaje y periodicidad trimestral/anual", "1 hallazgo comparando solo unidad de medida y periodicidad (sin fórmula)",
        f"{len(div)} hallazgo(s); atributos comparados: {', '.join(campos_div)}",
        len(div) == 1 and campos_div == ["periodicidad", "unidad_medida"],
    ))
    inc = por_tipo["incompletitud_atributos"]
    faltantes = sorted(inc[0].atributos_comparados["campos_faltantes"]) if inc else []
    filas.append(_fila(
        "CP-08", "CU4 Completitud", "Verificación de completitud de la hoja de vida",
        "Indicador sin objetivo, fórmula, unidad, línea base, periodicidad, fuente ni responsable", "1 hallazgo de incompletitud que lista los atributos faltantes",
        f"{len(inc)} hallazgo(s); faltantes: {', '.join(faltantes)}",
        len(inc) == 1 and {"objetivo", "formula", "unidad_medida", "fuente", "responsable"} <= set(faltantes),
    ))
    met = por_tipo["problema_metodologico"]
    filas.append(_fila(
        "CP-09", "CU4 Estructura metodológica", "Heurística léxica de nombres redactados como actividad",
        "Indicador 'Cierre de procesos contractuales'", "1 hallazgo de problema metodológico",
        f"{len(met)} hallazgo(s)" + (f" sobre «{met[0].indicador_principal.nombre}»" if met else ""),
        len(met) == 1 and met[0].indicador_principal.nombre == "Cierre de procesos contractuales",
    ))
    con_hallazgo_limpio = [h for h in hallazgos if limpio.pk in (h.indicador_principal_id, h.indicador_relacionado_id)]
    filas.append(_fila(
        "CP-10", "CU3-CU4 Caso negativo", "Todas las reglas",
        "Indicador sin inconsistencia esperada", "Ningún hallazgo asociado",
        f"{len(con_hallazgo_limpio)} hallazgo(s)", not con_hallazgo_limpio,
    ))
    filas.append(_fila(
        "CP-11", "CU3-CU4 Cobertura de reglas", "Cinco reglas de auditoría",
        "Conjunto de control completo", "Exactamente 5 hallazgos, uno por tipo, sin señalamientos adicionales ni ausentes",
        f"{len(hallazgos)} hallazgos: {', '.join(_tipos(ejecucion))}",
        len(hallazgos) == 5 and all(len(v) == 1 for v in por_tipo.values()),
    ))
    trazables = sum(1 for h in hallazgos if es_trazable(h))
    filas.append(_fila(
        "CP-12", "CU4 Trazabilidad", "Regla + atributos comparados + evidencia",
        "Los hallazgos de la ejecución", "100 % de hallazgos con regla, atributos y evidencia",
        f"{trazables} de {len(hallazgos)} trazables", trazables == len(hallazgos) > 0,
    ))
    filas.append(_fila(
        "CP-13", "CU5 Revisión humana", "Estado inicial del hallazgo",
        "Hallazgos recién generados", "Todos en estado pendiente",
        f"Estados: {sorted(set(h.estado for h in hallazgos))}", all(h.estado == "pendiente" for h in hallazgos),
    ))

    h = hallazgos[0]
    h.estado, h.revisado_por, h.revisado_en = "confirmado", usuario, timezone.now()
    h.save()
    h.refresh_from_db()
    filas.append(_fila(
        "CP-14", "CU5 Revisión humana", "Confirmación por el auditor",
        "Confirmar un hallazgo con un usuario autenticado", "Estado confirmado con revisor y fecha registrados",
        f"Estado {h.estado}; revisor {h.revisado_por}; fecha registrada: {h.revisado_en is not None}",
        h.estado == "confirmado" and h.revisado_por_id == usuario.pk and h.revisado_en is not None,
    ))

    from apps.reportes.views import exportar_excel
    solicitud = RequestFactory(SERVER_NAME="localhost").get("/reportes/hallazgos/excel/")
    solicitud.user = usuario
    respuesta = exportar_excel(solicitud)
    libro = load_workbook(io.BytesIO(respuesta.content))
    hoja = libro.active
    encabezados = [c.value for c in hoja[1]]
    evidencias = [fila[encabezados.index("Evidencia")] for fila in hoja.iter_rows(min_row=2, values_only=True)]
    de_la_ejecucion = {h.evidencia for h in hallazgos}
    filas.append(_fila(
        "CP-15", "CU5 Exportación", "Exportación a hoja de cálculo",
        "Exportar el listado de hallazgos", "Archivo .xlsx con una fila por hallazgo y su evidencia completa",
        f"{hoja.max_row - 1} filas; evidencias completas: {all(evidencias)}; hallazgos de la ejecución incluidos: {de_la_ejecucion <= set(evidencias)}",
        hoja.max_row - 1 >= len(hallazgos) and all(evidencias) and de_la_ejecucion <= set(evidencias),
    ))

    m = ejecucion.resumen_metricas
    filas.append(_fila(
        "CP-16", "CU5 Variables observables", "Seis variables observables",
        "Resumen de la ejecución", "Parámetros 9 de 11 (1 parcial), evidencia trazable 100 %, ejecución local confirmada, correspondencia disponible",
        f"Parámetros {m['parametros_operacionalizados']} de {m['parametros_totales']} ({m['parametros_parciales']} parcial); "
        f"evidencia {m['proporcion_evidencia_trazable']}; local {m['ejecucion_100_por_ciento_local']}; correspondencia {m['correspondencia_manual']}",
        m["parametros_operacionalizados"] == 9 and m["parametros_totales"] == 11 and m["parametros_parciales"] == 1
        and m["proporcion_evidencia_trazable"] == 1.0 and m["ejecucion_100_por_ciento_local"] is True
        and set(m["correspondencia_manual"]) == {"coincidente", "no_coincidente", "adicional"},
    ))

    if usar_modelo:
        segunda = ejecutar_auditoria(VIGENCIA, usuario=usuario)
    else:
        with mock.patch.object(nlp, "generar_embeddings_pendientes", lambda: 0), \
                mock.patch.object(nlp, "generar_pares_candidatos", _pares_simulados), _sin_red():
            segunda = ejecutar_auditoria(VIGENCIA, usuario=usuario)
    filas.append(_fila(
        "CP-17", "CU2-CU4 Repetibilidad", "Ejecución repetida",
        "Segunda auditoría sobre la misma vigencia", "Mismo resultado: 5 hallazgos con la misma distribución por tipo",
        f"{segunda.hallazgos.count()} hallazgos: {', '.join(_tipos(segunda))}",
        _tipos(segunda) == _tipos(ejecucion),
    ))

    from apps.dashboard.views import hallazgos as vista_hallazgos
    anonima = RequestFactory(SERVER_NAME="localhost").get("/hallazgos/")
    anonima.user = AnonymousUser()
    resp = vista_hallazgos(anonima)
    filas.append(_fila(
        "CP-18", "Seguridad de acceso", "Autenticación obligatoria",
        "Solicitud al listado de hallazgos sin sesión", "Redirección a la página de ingreso",
        f"Código {resp.status_code}; destino {resp.url}", resp.status_code == 302 and resp.url.startswith(reverse("login")),
    ))

    filas.append(_fila(
        "CP-19", "Operabilidad local", "Ejecución sin conexión externa",
        "Auditoría completa con la red bloqueada" if not usar_modelo else "Auditoría completa (la descarga inicial del modelo requiere red una sola vez)",
        "Sin conexiones salientes durante la ejecución; bandera de ejecución local registrada",
        ("Sin conexiones salientes; bandera " if not usar_modelo else "Bandera ") + str(ejecucion.ejecucion_100_por_ciento_local),
        ejecucion.ejecucion_100_por_ciento_local is True,
    ))
    return filas


def exportar_matriz_xlsx(filas: list[dict], ruta: str, usar_modelo: bool) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    libro = Workbook()
    hoja = libro.active
    hoja.title = "Matriz de validación"
    encabezados = ["Caso", "Módulo o funcionalidad", "Regla evaluada", "Datos de entrada", "Resultado esperado",
                   "Resultado obtenido", "Estado de cumplimiento", "Evidencia asociada"]
    hoja.append(encabezados)
    for c in hoja[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
    modo = "modelo de embeddings real" if usar_modelo else "pares candidatos simulados"
    evidencia = f"Salida de manage.py matriz_validacion ({modo}), {timezone.now():%Y-%m-%d %H:%M}"
    for f in filas:
        hoja.append([f["caso"], f["modulo"], f["regla"], f["entrada"], f["esperado"], f["obtenido"], f["estado"], evidencia])
    for fila in hoja.iter_rows(min_row=2):
        for c in fila:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        c = fila[6]
        c.fill = PatternFill("solid", fgColor="E2EFDA" if c.value == "Cumple" else "F8CBAD")
    for letra, ancho in zip("ABCDEFGH", (8, 26, 38, 44, 48, 52, 14, 40)):
        hoja.column_dimensions[letra].width = ancho
    hoja.freeze_panes = "A2"
    libro.save(ruta)
