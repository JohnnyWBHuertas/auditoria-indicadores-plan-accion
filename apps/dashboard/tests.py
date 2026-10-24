import io
from unittest import mock

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse

from apps.auditoria import validacion
from apps.auditoria.models import CorrespondenciaManual, Hallazgo
from apps.auditoria.reglas import ejecutar_auditoria
from apps.nlp_engine import embeddings as nlp


def usuario_con_rol(nombre, rol="Auditor"):
    """Crea un usuario y lo asigna a uno de los roles (grupos) del sistema."""
    call_command("seed_roles", stdout=io.StringIO())
    usuario = get_user_model().objects.create_user(nombre, password="x")
    usuario.groups.add(Group.objects.get(name=rol))
    return usuario


class TableroTest(TestCase):
    def setUp(self):
        call_command("seed_catalogo_gobernanza", stdout=io.StringIO())
        call_command("seed_datos_demo", stdout=io.StringIO())
        self.usuario = usuario_con_rol("auditor1")
        with mock.patch.object(nlp, "generar_embeddings_pendientes", lambda: 0), \
                mock.patch.object(nlp, "generar_pares_candidatos", validacion._pares_simulados):
            self.ejecucion = ejecutar_auditoria("2026", usuario=self.usuario)

    def test_requiere_sesion(self):
        for nombre in ("dashboard:inicio", "dashboard:hallazgos", "ingesta:cargar", "ingesta:historial", "reportes:exportar_excel"):
            respuesta = self.client.get(reverse(nombre))
            self.assertEqual(respuesta.status_code, 302, nombre)
            self.assertIn(reverse("login"), respuesta.url)

    def test_porcentajes_se_muestran_como_porcentaje(self):
        self.client.force_login(self.usuario)
        contenido = self.client.get(reverse("dashboard:inicio")).content.decode()
        self.assertIn("100%", contenido)  # evidencia trazable
        self.assertIn("9 / 11", contenido)  # parametros operacionalizados

    def test_tablero_sin_texto_de_llm(self):
        self.client.force_login(self.usuario)
        contenido = self.client.get(reverse("dashboard:hallazgos")).content.decode()
        self.assertNotIn("Qwen", contenido)

    def test_filtro_por_dependencia(self):
        self.client.force_login(self.usuario)
        h = Hallazgo.objects.filter(tipo="clonacion").get()
        dependencia = h.indicador_principal.dependencia_id
        con = self.client.get(reverse("dashboard:hallazgos"), {"dependencia": dependencia}).context["hallazgos"]
        self.assertTrue(all(x.indicador_principal.dependencia_id == dependencia for x in con))
        self.assertLess(len(con), Hallazgo.objects.count())

    def test_tarjetas_del_tablero_enlazan_al_listado_filtrado(self):
        self.client.force_login(self.usuario)
        contenido = self.client.get(reverse("dashboard:inicio")).content.decode()
        self.assertIn(f"?tipo=duplicidad&amp;ejecucion={self.ejecucion.pk}".replace("&amp;", "&"), contenido)
        self.assertIn("dependencia=", contenido)

    def test_filtro_por_tipo_y_ejecucion(self):
        self.client.force_login(self.usuario)
        r = self.client.get(reverse("dashboard:hallazgos"), {"tipo": "duplicidad", "ejecucion": self.ejecucion.pk})
        self.assertEqual(r.context["total"], 1)
        self.assertTrue(all(h.tipo == "duplicidad" for h in r.context["hallazgos"]))
        self.assertEqual(r.context["ejecucion_seleccionada"], self.ejecucion)

    def test_exportacion_respeta_el_filtro_de_tipo_y_ejecucion(self):
        self.client.force_login(self.usuario)
        from openpyxl import load_workbook
        r = self.client.get(reverse("reportes:exportar_excel"), {"tipo": "clonacion", "ejecucion": self.ejecucion.pk})
        self.assertEqual(load_workbook(io.BytesIO(r.content)).active.max_row, 2)  # encabezado + 1 hallazgo

    def test_confirmar_registra_revisor_y_fecha(self):
        self.client.force_login(self.usuario)
        h = Hallazgo.objects.first()
        self.client.post(reverse("dashboard:hallazgo_detalle", args=[h.pk]), {"accion": "confirmar", "observacion": "ok"})
        h.refresh_from_db()
        self.assertEqual((h.estado, h.revisado_por, h.observacion_revision), ("confirmado", self.usuario, "ok"))
        self.assertIsNotNone(h.revisado_en)

    def test_registrar_correspondencia_actualiza_la_variable_3(self):
        self.client.force_login(self.usuario)
        h = Hallazgo.objects.first()
        self.client.post(
            reverse("dashboard:hallazgo_detalle", args=[h.pk]),
            {"accion": "correspondencia", "tipo_correspondencia": "coincidente", "observacion_correspondencia": "igual"},
        )
        self.assertEqual(CorrespondenciaManual.objects.get().tipo_correspondencia, "coincidente")
        self.ejecucion.refresh_from_db()
        self.assertEqual(self.ejecucion.resumen_metricas["correspondencia_manual"]["coincidente"], 1)

    def test_exportacion_incluye_dependencia_y_no_columna_llm(self):
        self.client.force_login(self.usuario)
        respuesta = self.client.get(reverse("reportes:exportar_excel"))
        self.assertEqual(respuesta.status_code, 200)
        from openpyxl import load_workbook
        encabezados = [c.value for c in load_workbook(io.BytesIO(respuesta.content)).active[1]]
        self.assertIn("Dependencia principal", encabezados)
        self.assertFalse(any("Qwen" in str(e) for e in encabezados))


class RolesTest(TestCase):
    def test_grupos_y_permisos(self):
        call_command("seed_roles", usuarios=True, password="clave-de-prueba", stdout=io.StringIO())
        Usuario = get_user_model()
        auditor, administrador = Usuario.objects.get(username="auditor"), Usuario.objects.get(username="administrador")
        self.assertFalse(auditor.is_staff)
        self.assertTrue(administrador.is_staff)
        self.assertTrue(auditor.has_perm("auditoria.change_hallazgo"))
        self.assertFalse(auditor.has_perm("auditoria.change_reglaauditoria"))
        self.assertTrue(administrador.has_perm("auditoria.change_reglaauditoria"))
        self.assertTrue(administrador.has_perm("ingesta.add_plantillamapeocolumnas"))


class DetalleLegibleTest(TestCase):
    """Los atributos comparados se muestran en lenguaje comprensible, no como JSON."""

    def setUp(self):
        call_command("seed_catalogo_gobernanza", stdout=io.StringIO())
        call_command("seed_datos_demo", stdout=io.StringIO())
        self.usuario = usuario_con_rol("auditor2")
        with mock.patch.object(nlp, "generar_embeddings_pendientes", lambda: 0), \
                mock.patch.object(nlp, "generar_pares_candidatos", validacion._pares_simulados):
            ejecutar_auditoria("2026", usuario=self.usuario)

    def _detalle(self, tipo):
        from apps.auditoria.presentacion import presentar
        return presentar(Hallazgo.objects.get(tipo=tipo))

    def test_divergencia_compara_unidad_y_periodicidad(self):
        d = self._detalle("divergencia_operativa")
        filas = {f["etiqueta"]: f for f in d["filas"]}
        self.assertEqual((filas["Unidad de medida"]["a"], filas["Unidad de medida"]["b"]), ("Número", "Porcentaje"))
        self.assertEqual((filas["Periodicidad"]["a"], filas["Periodicidad"]["b"]), ("Trimestral", "Anual"))
        self.assertFalse(filas["Periodicidad"]["igual"])
        self.assertEqual(d["medidas"][0][1], "90 %")

    def test_clonacion_muestra_nombres_de_dependencia_y_marca_diferencias(self):
        d = self._detalle("clonacion")
        filas = {f["etiqueta"]: f for f in d["filas"]}
        self.assertTrue(filas["Nombre del indicador"]["igual"])
        self.assertFalse(filas["Producto o actividad asociado"]["igual"])
        self.assertIn("Grupo de Convocatorias", filas["Dependencia"]["a"])  # nombre, no un numero

    def test_incompletitud_lista_lo_que_falta(self):
        d = self._detalle("incompletitud_atributos")
        self.assertIn("Fórmula de cálculo", d["faltantes"])
        self.assertIn("Nombre del indicador", d["diligenciados"])
        self.assertIn("Faltan", d["explicacion"])

    def test_la_pagina_no_muestra_el_json_fuera_del_detalle_tecnico(self):
        self.client.force_login(self.usuario)
        for h in Hallazgo.objects.all():
            html = self.client.get(reverse("dashboard:hallazgo_detalle", args=[h.pk])).content.decode()
            visible = html.split("<details")[0]
            self.assertNotIn("'indicador_a'", visible)
            self.assertIn("Qué se comparó", visible)


class DetalleDeVariablesTest(TestCase):
    """Cada tarjeta del tablero lleva a una pagina de detalle con su informacion."""

    def setUp(self):
        call_command("seed_catalogo_gobernanza", stdout=io.StringIO())
        call_command("seed_datos_demo", stdout=io.StringIO())
        self.usuario = usuario_con_rol("auditor3")
        with mock.patch.object(nlp, "generar_embeddings_pendientes", lambda: 0), \
                mock.patch.object(nlp, "generar_pares_candidatos", validacion._pares_simulados):
            self.ejecucion = ejecutar_auditoria("2026", usuario=self.usuario)
        self.client.force_login(self.usuario)

    def _get(self, nombre, **parametros):
        return self.client.get(reverse(nombre, args=[self.ejecucion.pk]), parametros)

    def test_las_paginas_exigen_sesion(self):
        self.client.logout()
        for nombre in ("parametros", "correspondencia", "trazabilidad", "completitud"):
            r = self.client.get(reverse(f"dashboard:detalle_{nombre}", args=[self.ejecucion.pk]))
            self.assertEqual(r.status_code, 302, nombre)

    def test_tablero_enlaza_a_los_cuatro_detalles(self):
        html = self.client.get(reverse("dashboard:inicio")).content.decode()
        for nombre in ("parametros", "correspondencia", "trazabilidad", "completitud"):
            self.assertIn(reverse(f"dashboard:detalle_{nombre}", args=[self.ejecucion.pk]), html)

    def test_matriz_de_parametros(self):
        r = self._get("dashboard:detalle_parametros")
        self.assertEqual((r.context["si"], r.context["parcial"], r.context["no"], r.context["total"]), (9, 1, 1, 11))
        filas = {f["parametro"].codigo: f for f in r.context["filas"]}
        self.assertEqual({x.tipo_hallazgo for x in filas["P04"]["reglas"]}, {"duplicidad", "clonacion"})
        self.assertEqual(filas["P04"]["hallazgos"], 2)  # un hallazgo de duplicidad y uno de clonacion
        self.assertIsNone(filas["P01"]["hallazgos"])    # atributo de datos, sin regla propia
        self.assertIsNone(filas["P09"]["hallazgos"])    # principio de diseño
        self.assertContains(r, "Atributo o principio de diseño")
        self.assertContains(r, "Sin regla: no operacionalizado")  # P03
        self.assertContains(r, "cubierto en parte")                # P02

    def test_correspondencia_clasifica_por_avance(self):
        h = Hallazgo.objects.first()
        CorrespondenciaManual.objects.create(hallazgo=h, tipo_correspondencia="coincidente")
        r = self._get("dashboard:detalle_correspondencia")
        cuentas = {clave: n for clave, _, n in r.context["grupos"]}
        self.assertEqual(cuentas, {"sin_registrar": 4, "coincidente": 1, "no_coincidente": 0, "adicional": 0})
        self.assertEqual(r.context["total_grupo"], 4)  # por defecto, los sin registrar
        r = self._get("dashboard:detalle_correspondencia", grupo="coincidente")
        self.assertEqual([x.pk for x in r.context["hallazgos"]], [h.pk])

    def test_trazabilidad_sin_excepciones(self):
        r = self._get("dashboard:detalle_trazabilidad")
        self.assertEqual((r.context["total"], r.context["trazables"]), (5, 5))
        self.assertEqual(r.context["excepciones"], [])
        self.assertContains(r, "Ninguna")

    def test_completitud_por_atributo_dependencia_y_lista(self):
        r = self._get("dashboard:detalle_completitud")
        atributos = {a["campo"]: a for a in r.context["atributos"]}
        self.assertEqual(atributos["nombre"]["pct"], 100)
        self.assertLess(atributos["fuente"]["pct"], 100)
        self.assertEqual(r.context["atributos"][0]["pct"], min(a["pct"] for a in r.context["atributos"]))  # peores primero
        self.assertTrue(r.context["dependencias"])
        r = self._get("dashboard:detalle_completitud", atributo="fuente")
        self.assertEqual(r.context["filtro"], "fuente")
        self.assertTrue(all("Fuente de verificación" in f["faltan"] for f in r.context["lista"]))


class CompararVigenciasTest(TestCase):
    """Comparacion entre vigencias: indicadores, hallazgos y variacion."""

    def setUp(self):
        from apps.gobernanza.models import Dependencia, EjecucionAuditoria
        from apps.ingesta.models import Indicador

        self.usuario = usuario_con_rol("auditor4")
        self.client.force_login(self.usuario)
        dep = Dependencia.objects.create(nombre="Dep")
        call_command("seed_catalogo_gobernanza", stdout=io.StringIO())
        regla = Hallazgo._meta.get_field("regla_aplicada").related_model.objects.get(tipo_hallazgo="duplicidad")
        regla2 = Hallazgo._meta.get_field("regla_aplicada").related_model.objects.get(tipo_hallazgo="clonacion")

        def vigencia(anio, n_indicadores, tipos):
            indicadores = [Indicador.objects.create(nombre=f"I{anio}-{i}", dependencia=dep, vigencia=anio) for i in range(n_indicadores)]
            ejecucion = EjecucionAuditoria.objects.create(vigencia=anio)
            for tipo in tipos:
                Hallazgo.objects.create(
                    tipo=tipo, ejecucion=ejecucion, regla_aplicada=regla if tipo == "duplicidad" else regla2,
                    indicador_principal=indicadores[0], atributos_comparados={"x": 1}, evidencia="e",
                )
            return ejecucion

        vigencia("2024", 10, ["duplicidad"] * 4)                       # 4 hallazgos
        vigencia("2025", 20, ["duplicidad"] * 4 + ["clonacion"] * 2)    # 6 hallazgos: +50 %
        vigencia("2026", 5, [])                                        # 0 hallazgos: -100 %
        Indicador.objects.create(nombre="Sin auditoria", dependencia=dep, vigencia="2027")

    def _get(self, *vigencias):
        return self.client.get(reverse("dashboard:comparar"), {"vigencias": list(vigencias)})

    def test_requiere_sesion(self):
        self.client.logout()
        self.assertEqual(self.client.get(reverse("dashboard:comparar")).status_code, 302)

    def test_sin_seleccion_muestra_todas_las_vigencias(self):
        r = self.client.get(reverse("dashboard:comparar"))
        self.assertEqual([f["vigencia"] for f in r.context["filas"]], ["2024", "2025", "2026", "2027"])
        self.assertTrue(r.context["sin_seleccion"])

    def test_seleccion_multiple_cuenta_indicadores_y_hallazgos(self):
        r = self._get("2024", "2025")
        filas = {f["vigencia"]: f for f in r.context["filas"]}
        self.assertEqual(list(filas), ["2024", "2025"])
        self.assertEqual((filas["2024"]["indicadores"], filas["2024"]["hallazgos"]), (10, 4))
        self.assertEqual((filas["2025"]["indicadores"], filas["2025"]["hallazgos"]), (20, 6))
        self.assertEqual(r.context["total_indicadores"], 30)

    def test_variacion_porcentual_respecto_de_la_anterior(self):
        r = self._get("2024", "2025", "2026")
        v = [f["variacion"] for f in r.context["filas"]]
        self.assertEqual(v, [None, 50.0, -100.0])

    def test_variacion_salta_vigencias_sin_auditoria(self):
        r = self._get("2025", "2027")
        self.assertEqual([f["variacion"] for f in r.context["filas"]], [None, None])
        self.assertEqual(r.context["sin_auditoria"], ["2027"])

    def test_variacion_no_se_define_sin_base(self):
        from apps.dashboard.views_comparar import variacion
        self.assertIsNone(variacion(5, 0))
        self.assertIsNone(variacion(None, 3))
        self.assertEqual(variacion(3, 4), -25.0)

    def test_total_de_indicadores_incluye_los_que_no_tienen_hallazgos(self):
        r = self._get("2025", "2027")
        filas = {f["vigencia"]: f for f in r.context["filas"]}
        self.assertEqual(filas["2025"]["indicadores"], 20)      # todos los de la vigencia
        self.assertEqual(filas["2025"]["involucrados"], 1)      # solo uno aparece en hallazgos
        self.assertEqual(filas["2025"]["pct_involucrados"], 5)  # 1 / 20
        self.assertEqual(filas["2027"]["indicadores"], 1)
        self.assertIsNone(filas["2027"]["involucrados"])        # sin auditoria
        self.assertContains(r, "Total de indicadores")
        self.assertContains(r, "Indicadores con hallazgos")

    def test_involucrados_cuenta_principal_y_relacionado_sin_repetir(self):
        from apps.gobernanza.models import EjecucionAuditoria
        from apps.ingesta.models import Indicador
        a, b = Indicador.objects.filter(vigencia="2024")[:2]
        ejecucion = EjecucionAuditoria.objects.create(vigencia="2024")
        regla = Hallazgo._meta.get_field("regla_aplicada").related_model.objects.get(tipo_hallazgo="duplicidad")
        for _ in range(2):  # dos hallazgos sobre el mismo par
            Hallazgo.objects.create(tipo="duplicidad", ejecucion=ejecucion, regla_aplicada=regla, indicador_principal=a,
                                    indicador_relacionado=b, atributos_comparados={"x": 1}, evidencia="e")
        self.assertEqual(self._get("2024").context["filas"][0]["involucrados"], 2)

    def test_hallazgos_por_indicador_y_por_tipo(self):
        r = self._get("2025")
        fila = r.context["filas"][0]
        self.assertEqual(fila["por_indicador"], 0.3)  # 6 / 20
        self.assertEqual(fila["por_tipo"], {"duplicidad": 4, "clonacion": 2})

    def test_datos_de_la_grafica(self):
        r = self._get("2024", "2025")
        self.assertEqual(r.context["grafica"], {
            "etiquetas": ["2024", "2025"], "indicadores": [10, 20], "hallazgos": [4, 6], "variacion": [None, 50.0],
        })
        self.assertContains(r, 'id="datos-grafica"')

    def test_usa_la_ultima_ejecucion_de_cada_vigencia(self):
        from apps.gobernanza.models import EjecucionAuditoria
        nueva = EjecucionAuditoria.objects.create(vigencia="2024")  # mas reciente y sin hallazgos
        r = self._get("2024")
        self.assertEqual(r.context["filas"][0]["ejecucion"], nueva)
        self.assertEqual(r.context["filas"][0]["hallazgos"], 0)

    def test_ignora_vigencias_inexistentes(self):
        r = self._get("2025", "1999")
        self.assertEqual([f["vigencia"] for f in r.context["filas"]], ["2025"])


class ControlDeAccesoPorRolTest(TestCase):
    """RBAC: cada rol (grupo) solo accede a las funcionalidades de sus permisos."""

    def setUp(self):
        call_command("seed_catalogo_gobernanza", stdout=io.StringIO())
        call_command("seed_datos_demo", stdout=io.StringIO())
        self.auditor = usuario_con_rol("rol_auditor", "Auditor")
        self.consulta = usuario_con_rol("rol_consulta", "Consulta")
        self.admin = usuario_con_rol("rol_admin", "Administrador")
        with mock.patch.object(nlp, "generar_embeddings_pendientes", lambda: 0), \
                mock.patch.object(nlp, "generar_pares_candidatos", validacion._pares_simulados):
            self.ejecucion = ejecutar_auditoria("2026", usuario=self.auditor)
        self.hallazgo = Hallazgo.objects.filter(ejecucion=self.ejecucion).first()

    def test_consulta_puede_ver_pero_no_operar(self):
        self.client.force_login(self.consulta)
        for nombre in ("dashboard:inicio", "dashboard:hallazgos", "dashboard:comparar", "ingesta:historial", "reportes:exportar_excel"):
            self.assertEqual(self.client.get(reverse(nombre)).status_code, 200, nombre)
        self.assertEqual(self.client.get(reverse("ingesta:cargar")).status_code, 403)
        self.assertEqual(self.client.post(reverse("dashboard:inicio"), {"accion": "ejecutar_auditoria", "vigencia": "2026"}).status_code, 403)
        url = reverse("dashboard:hallazgo_detalle", args=[self.hallazgo.pk])
        self.assertEqual(self.client.post(url, {"accion": "confirmar"}).status_code, 403)
        self.hallazgo.refresh_from_db()
        self.assertEqual(self.hallazgo.estado, "pendiente")

    def test_auditor_carga_audita_y_revisa(self):
        self.client.force_login(self.auditor)
        self.assertEqual(self.client.get(reverse("ingesta:cargar")).status_code, 200)
        url = reverse("dashboard:hallazgo_detalle", args=[self.hallazgo.pk])
        self.assertEqual(self.client.post(url, {"accion": "confirmar", "observacion": "ok"}).status_code, 302)
        self.hallazgo.refresh_from_db()
        self.assertEqual(self.hallazgo.estado, "confirmado")
        self.assertEqual(self.hallazgo.revisado_por, self.auditor)

    def test_solo_el_administrador_entra_a_la_administracion(self):
        self.admin.is_staff = True
        self.admin.save()
        self.client.force_login(self.auditor)
        self.assertEqual(self.client.get("/admin/").status_code, 302)  # no es personal del sitio
        self.client.force_login(self.admin)
        self.assertEqual(self.client.get("/admin/").status_code, 200)

    def test_pagina_403_es_comprensible(self):
        self.client.force_login(self.consulta)
        respuesta = self.client.get(reverse("ingesta:cargar"))
        self.assertContains(respuesta, "No tiene permiso para esta acción", status_code=403)
